"""Correct the synthetic HC sales file (code review 2026-10, finding C58).

The first version applied VAT twice on every sale row: 'Գումար (ԱԱՀ-ով)' was 1.2 x (price x qty - discount)
instead of (price x qty - discount), while 'Գումար (առանց ԱԱՀ)' and the return rows were right. This script
divides that column by 1.2 on sale rows only and checks the result against the ex-VAT column.

The column stays an amount AFTER discount (finding C57): map it as net_amount, not gross_amount.

Usage: python3 -I fix_synthetic_hc_vat.py <file.xlsx> [<file.xlsx> ...]   (files are rewritten in place)
"""
import sys

import openpyxl

INC, EX = 'Գումար (ԱԱՀ-ով)', 'Գումար (առանց ԱԱՀ)'
OP, SALE_PREFIX = 'Գործառնության տեսակ', 'Վաճառք'


def fix(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    inc, ex, op = header.index(INC) + 1, header.index(EX) + 1, header.index(OP) + 1
    fixed = 0
    for r in range(2, ws.max_row + 1):
        if not str(ws.cell(r, op).value).startswith(SALE_PREFIX):
            continue
        old = ws.cell(r, inc).value
        new = round(old / 1.2)
        if abs(new / 1.2 - ws.cell(r, ex).value) > 1:
            raise SystemExit(f'{path} row {r}: {old} is not 1.2 x VAT-inclusive amount; already fixed?')
        ws.cell(r, inc).value = new
        fixed += 1
    wb.save(path)
    print(f'{path}: {fixed} sale rows corrected')


if __name__ == '__main__':
    for p in sys.argv[1:]:
        fix(p)
